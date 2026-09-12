from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form
from fastapi.responses import StreamingResponse
from io import BytesIO
import openpyxl
from sqlalchemy.orm import Session
from typing import List
from datetime import datetime, timezone
import random

import models
import schemas
from auth import get_db, require_admin, get_current_user

router = APIRouter()


# ---------- SUBJECT ----------

@router.post("/subjects", response_model=schemas.SubjectResponse)
def create_subject(
    subject: schemas.SubjectCreate,
    db: Session = Depends(get_db),
    admin: models.User = Depends(require_admin),
):
    new_subject = models.Subject(name=subject.name, code=subject.code)
    db.add(new_subject)
    db.commit()
    db.refresh(new_subject)
    return new_subject


@router.get("/subjects", response_model=List[schemas.SubjectResponse])
def list_subjects(db: Session = Depends(get_db), current_user: models.User = Depends(get_current_user)):
    return db.query(models.Subject).all()


# ---------- CHAPTER ----------

@router.post("/chapters", response_model=schemas.ChapterResponse)
def create_chapter(
    chapter: schemas.ChapterCreate,
    db: Session = Depends(get_db),
    admin: models.User = Depends(require_admin),
):
    subject = db.query(models.Subject).filter(models.Subject.id == chapter.subject_id).first()
    if not subject:
        raise HTTPException(status_code=404, detail="Môn học không tồn tại")

    new_chapter = models.Chapter(name=chapter.name, subject_id=chapter.subject_id)
    db.add(new_chapter)
    db.commit()
    db.refresh(new_chapter)
    return new_chapter


@router.get("/chapters", response_model=List[schemas.ChapterResponse])
def list_chapters(
    subject_id: int = None,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    query = db.query(models.Chapter)
    if subject_id:
        query = query.filter(models.Chapter.subject_id == subject_id)
    return query.all()


# ---------- QUESTION ----------

@router.post("/questions", response_model=schemas.QuestionResponse)
def create_question(
    question: schemas.QuestionCreate,
    db: Session = Depends(get_db),
    admin: models.User = Depends(require_admin),
):
    chapter = db.query(models.Chapter).filter(models.Chapter.id == question.chapter_id).first()
    if not chapter:
        raise HTTPException(status_code=404, detail="Chương học không tồn tại")

    if len(question.answers) < 2:
        raise HTTPException(status_code=400, detail="Câu hỏi cần ít nhất 2 đáp án")

    if not any(a.is_correct for a in question.answers):
        raise HTTPException(status_code=400, detail="Phải có ít nhất 1 đáp án đúng")

    new_question = models.Question(
        content=question.content,
        difficulty=question.difficulty,
        chapter_id=question.chapter_id,
    )
    db.add(new_question)
    db.commit()
    db.refresh(new_question)

    for a in question.answers:
        new_answer = models.Answer(
            content=a.content,
            is_correct=a.is_correct,
            question_id=new_question.id,
        )
        db.add(new_answer)

    db.commit()
    db.refresh(new_question)
    return new_question


@router.get("/questions", response_model=List[schemas.QuestionResponse])
def list_questions(
    chapter_id: int = None,
    db: Session = Depends(get_db),
    admin: models.User = Depends(require_admin),
):
    query = db.query(models.Question)
    if chapter_id:
        query = query.filter(models.Question.chapter_id == chapter_id)
    return query.all()


@router.delete("/questions/{question_id}")
def delete_question(
    question_id: int,
    db: Session = Depends(get_db),
    admin: models.User = Depends(require_admin),
):
    question = db.query(models.Question).filter(models.Question.id == question_id).first()
    if not question:
        raise HTTPException(status_code=404, detail="Câu hỏi không tồn tại")

    used_in_exam = db.query(models.ExamQuestion).filter(models.ExamQuestion.question_id == question_id).first()
    if used_in_exam:
        raise HTTPException(
            status_code=400,
            detail="Câu hỏi đã được sử dụng trong đề thi, không thể xóa. Hãy xóa hoặc ẩn đề thi liên quan trước.",
        )

    db.delete(question)
    db.commit()
    return {"message": "Đã xóa câu hỏi"}

@router.get("/questions/import-template")
def download_import_template(admin: models.User = Depends(require_admin)):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Câu hỏi"
    ws.append(["Câu hỏi", "Đáp án A", "Đáp án B", "Đáp án C", "Đáp án D", "Đáp án đúng (A/B/C/D)", "Độ khó (easy/medium/hard)"])
    ws.append([
        "Thủ đô của Việt Nam là gì?",
        "Hà Nội", "Hồ Chí Minh", "Đà Nẵng", "Huế",
        "A", "easy",
    ])

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)

    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=mau_import_cau_hoi.xlsx"},
    )


@router.post("/questions/import", response_model=schemas.ImportResultResponse)
async def import_questions(
    chapter_id: int = Form(...),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    admin: models.User = Depends(require_admin),
):
    chapter = db.query(models.Chapter).filter(models.Chapter.id == chapter_id).first()
    if not chapter:
        raise HTTPException(status_code=404, detail="Chương học không tồn tại")

    if not file.filename.lower().endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="Chỉ hỗ trợ file Excel (.xlsx)")

    content = await file.read()
    try:
        wb = openpyxl.load_workbook(BytesIO(content))
        ws = wb.active
    except Exception:
        raise HTTPException(status_code=400, detail="Không đọc được file, vui lòng dùng đúng file mẫu")

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    success_count = 0
    errors = []

    for idx, row in enumerate(rows, start=2):
        if row is None or all(cell is None for cell in row):
            continue

        padded = (list(row) + [None] * 7)[:7]
        content_text, a, b, c, d, correct_letter, difficulty = padded

        if not content_text or not str(content_text).strip():
            errors.append(f"Dòng {idx}: thiếu nội dung câu hỏi")
            continue

        options = {"A": a, "B": b, "C": c, "D": d}
        filled_options = {k: v for k, v in options.items() if v and str(v).strip()}

        if len(filled_options) < 2:
            errors.append(f"Dòng {idx}: cần ít nhất 2 đáp án")
            continue

        correct_letter = str(correct_letter).strip().upper() if correct_letter else ""
        if correct_letter not in filled_options:
            errors.append(f"Dòng {idx}: đáp án đúng '{correct_letter}' không hợp lệ hoặc để trống")
            continue

        diff = str(difficulty).strip().lower() if difficulty else "medium"
        if diff not in ("easy", "medium", "hard"):
            diff = "medium"

        new_question = models.Question(content=str(content_text).strip(), difficulty=diff, chapter_id=chapter_id)
        db.add(new_question)
        db.commit()
        db.refresh(new_question)

        for letter, text in filled_options.items():
            db.add(models.Answer(
                content=str(text).strip(),
                is_correct=(letter == correct_letter),
                question_id=new_question.id,
            ))
        db.commit()
        success_count += 1

    return schemas.ImportResultResponse(success_count=success_count, error_count=len(errors), errors=errors)


# ---------- EXAM ----------

@router.post("/exams", response_model=schemas.ExamResponse)
def create_exam(
    exam: schemas.ExamCreate,
    db: Session = Depends(get_db),
    admin: models.User = Depends(require_admin),
):
    subject = db.query(models.Subject).filter(models.Subject.id == exam.subject_id).first()
    if not subject:
        raise HTTPException(status_code=404, detail="Môn học không tồn tại")

    if exam.chapter_id:
        chapter = db.query(models.Chapter).filter(
            models.Chapter.id == exam.chapter_id,
            models.Chapter.subject_id == exam.subject_id,
        ).first()
        if not chapter:
            raise HTTPException(status_code=404, detail="Chương không thuộc môn học đã chọn")
        chapter_ids = [exam.chapter_id]
    else:
        chapter_ids = [c.id for c in db.query(models.Chapter).filter(models.Chapter.subject_id == exam.subject_id)]

    query = db.query(models.Question).filter(models.Question.chapter_id.in_(chapter_ids))
    if exam.difficulty:
        query = query.filter(models.Question.difficulty == exam.difficulty)

    all_questions = query.all()

    if len(all_questions) < exam.question_count:
        raise HTTPException(
            status_code=400,
            detail=f"Ngân hàng câu hỏi chỉ có {len(all_questions)} câu, không đủ {exam.question_count} câu yêu cầu",
        )

    selected_questions = random.sample(all_questions, exam.question_count)

    new_exam = models.Exam(
        title=exam.title,
        subject_id=exam.subject_id,
        chapter_id=exam.chapter_id,
        question_count=exam.question_count,
        duration=exam.duration_minutes,
        created_by=admin.id,
    )
    db.add(new_exam)
    db.commit()
    db.refresh(new_exam)

    for q in selected_questions:
        db.add(models.ExamQuestion(exam_id=new_exam.id, question_id=q.id))
    db.commit()

    return schemas.ExamResponse(
        id=new_exam.id,
        title=new_exam.title,
        subject_id=new_exam.subject_id,
        chapter_id=new_exam.chapter_id,
        question_count=new_exam.question_count,
        duration_minutes=new_exam.duration,
        is_archived=new_exam.is_archived,
        attempts_count=0,
    )


@router.get("/exams", response_model=List[schemas.ExamResponse])
def list_exams(
    subject_id: int = None,
    chapter_id: int = None,
    subject_wide: bool = False,
    include_archived: bool = False,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    query = db.query(models.Exam)
    if subject_id:
        query = query.filter(models.Exam.subject_id == subject_id)
    if chapter_id:
        query = query.filter(models.Exam.chapter_id == chapter_id)
    elif subject_wide:
        query = query.filter(models.Exam.chapter_id.is_(None))

    if not include_archived:
        query = query.filter(models.Exam.is_archived == False)

    exams = query.all()
    result = []
    for e in exams:
        attempts_count = db.query(models.Attempt).filter(models.Attempt.exam_id == e.id).count()
        result.append(schemas.ExamResponse(
            id=e.id, title=e.title, subject_id=e.subject_id, chapter_id=e.chapter_id,
            question_count=e.question_count, duration_minutes=e.duration,
            is_archived=e.is_archived, attempts_count=attempts_count,
        ))
    return result


@router.get("/exams/{exam_id}", response_model=schemas.ExamDetailResponse)
def get_exam_detail(
    exam_id: int,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    exam = db.query(models.Exam).filter(models.Exam.id == exam_id).first()
    if not exam:
        raise HTTPException(status_code=404, detail="Đề thi không tồn tại")

    if exam.is_archived:
        raise HTTPException(status_code=400, detail="Đề thi này đã ngừng sử dụng")

    questions = [eq.question for eq in exam.exam_questions]

    question_responses = []
    for q in questions:
        answers = list(q.answers)
        random.shuffle(answers)
        question_responses.append(
            schemas.ExamQuestionResponse(id=q.id, content=q.content, answers=answers)
        )

    return schemas.ExamDetailResponse(
        id=exam.id,
        title=exam.title,
        duration_minutes=exam.duration,
        questions=question_responses,
    )


@router.delete("/exams/{exam_id}")
def delete_exam(
    exam_id: int,
    db: Session = Depends(get_db),
    admin: models.User = Depends(require_admin),
):
    exam = db.query(models.Exam).filter(models.Exam.id == exam_id).first()
    if not exam:
        raise HTTPException(status_code=404, detail="Đề thi không tồn tại")

    has_attempts = db.query(models.Attempt).filter(models.Attempt.exam_id == exam_id).first()
    if has_attempts:
        raise HTTPException(
            status_code=400,
            detail="Đề thi đã có sinh viên làm bài, không thể xóa. Hãy dùng chức năng Ẩn đề thi thay thế.",
        )

    db.query(models.ExamQuestion).filter(models.ExamQuestion.exam_id == exam_id).delete()
    db.delete(exam)
    db.commit()
    return {"message": "Đã xóa đề thi"}


@router.patch("/exams/{exam_id}/archive", response_model=schemas.ExamResponse)
def toggle_archive_exam(
    exam_id: int,
    db: Session = Depends(get_db),
    admin: models.User = Depends(require_admin),
):
    exam = db.query(models.Exam).filter(models.Exam.id == exam_id).first()
    if not exam:
        raise HTTPException(status_code=404, detail="Đề thi không tồn tại")

    exam.is_archived = not exam.is_archived
    db.commit()
    db.refresh(exam)

    attempts_count = db.query(models.Attempt).filter(models.Attempt.exam_id == exam_id).count()
    return schemas.ExamResponse(
        id=exam.id, title=exam.title, subject_id=exam.subject_id, chapter_id=exam.chapter_id,
        question_count=exam.question_count, duration_minutes=exam.duration,
        is_archived=exam.is_archived, attempts_count=attempts_count,
    )


# ---------- ATTEMPT (làm bài & chấm điểm) ----------

@router.post("/exams/{exam_id}/submit", response_model=schemas.AttemptResultResponse)
def submit_exam(
    exam_id: int,
    submission: schemas.AttemptSubmit,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    exam = db.query(models.Exam).filter(models.Exam.id == exam_id).first()
    if not exam:
        raise HTTPException(status_code=404, detail="Đề thi không tồn tại")

    valid_question_ids = {eq.question_id for eq in exam.exam_questions}

    new_attempt = models.Attempt(
        user_id=current_user.id,
        exam_id=exam_id,
        finished_at=datetime.now(timezone.utc),
    )
    db.add(new_attempt)
    db.commit()
    db.refresh(new_attempt)

    details = []
    correct_count = 0

    for item in submission.answers:
        if item.question_id not in valid_question_ids:
            continue

        question = db.query(models.Question).filter(models.Question.id == item.question_id).first()
        if not question:
            continue

        correct_answer = next((a for a in question.answers if a.is_correct), None)
        selected_answer = db.query(models.Answer).filter(models.Answer.id == item.selected_answer_id).first()

        is_correct = bool(selected_answer and correct_answer and selected_answer.id == correct_answer.id)
        if is_correct:
            correct_count += 1

        attempt_answer = models.AttemptAnswer(
            attempt_id=new_attempt.id,
            question_id=item.question_id,
            selected_answer_id=item.selected_answer_id,
            is_correct=is_correct,
        )
        db.add(attempt_answer)

        details.append(
            schemas.AttemptAnswerResult(
                question_id=question.id,
                question_content=question.content,
                selected_answer_id=item.selected_answer_id,
                selected_answer_content=selected_answer.content if selected_answer else "",
                is_correct=is_correct,
                correct_answer_id=correct_answer.id if correct_answer else 0,
                correct_answer_content=correct_answer.content if correct_answer else "",
    )
)

    total_questions = len(valid_question_ids)
    score = round((correct_count / total_questions) * 10, 2) if total_questions > 0 else 0

    new_attempt.score = score
    db.commit()

    return schemas.AttemptResultResponse(
        id=new_attempt.id,
        exam_id=exam_id,
        score=score,
        total_questions=total_questions,
        correct_count=correct_count,
        details=details,
    )


@router.get("/attempts/me", response_model=List[schemas.AttemptSummaryResponse])
def my_attempts(db: Session = Depends(get_db), current_user: models.User = Depends(get_current_user)):
    attempts = db.query(models.Attempt).filter(models.Attempt.user_id == current_user.id).order_by(models.Attempt.started_at.desc()).all()
    return [
        schemas.AttemptSummaryResponse(
            id=a.id, exam_id=a.exam_id, exam_title=a.exam.title,
            score=a.score, started_at=a.started_at, finished_at=a.finished_at,
        )
        for a in attempts
    ]

@router.get("/attempts/{attempt_id}", response_model=schemas.AttemptDetailResponse)
def get_attempt_detail(
    attempt_id: int,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    attempt = db.query(models.Attempt).filter(
        models.Attempt.id == attempt_id,
        models.Attempt.user_id == current_user.id,
    ).first()
    if not attempt:
        raise HTTPException(status_code=404, detail="Không tìm thấy lượt làm bài này")

    attempt_answers = db.query(models.AttemptAnswer).filter(models.AttemptAnswer.attempt_id == attempt_id).all()

    details = []
    correct_count = 0
    for aa in attempt_answers:
        question = aa.question
        selected_answer = aa.selected_answer
        correct_answer = next((a for a in question.answers if a.is_correct), None)
        if aa.is_correct:
            correct_count += 1

        details.append(schemas.AttemptAnswerResult(
            question_id=question.id,
            question_content=question.content,
            selected_answer_id=aa.selected_answer_id,
            selected_answer_content=selected_answer.content if selected_answer else "",
            is_correct=aa.is_correct,
            correct_answer_id=correct_answer.id if correct_answer else 0,
            correct_answer_content=correct_answer.content if correct_answer else "",
        ))

    return schemas.AttemptDetailResponse(
        id=attempt.id, exam_id=attempt.exam_id, exam_title=attempt.exam.title,
        score=attempt.score, total_questions=len(attempt_answers), correct_count=correct_count,
        started_at=attempt.started_at, finished_at=attempt.finished_at, details=details,
    )

# ---------- THỐNG KÊ ----------

@router.get("/stats/me", response_model=schemas.StudentStatsResponse)
def my_stats(db: Session = Depends(get_db), current_user: models.User = Depends(get_current_user)):
    attempts = db.query(models.Attempt).filter(models.Attempt.user_id == current_user.id).all()
    total_attempts = len(attempts)
    average_score = round(sum(a.score for a in attempts) / total_attempts, 2) if total_attempts > 0 else 0
    highest_score = max((a.score for a in attempts), default=0)

    attempt_ids = [a.id for a in attempts]
    attempt_answers = (
        db.query(models.AttemptAnswer)
        .filter(models.AttemptAnswer.attempt_id.in_(attempt_ids))
        .all()
    )

    chapter_stats = {}
    for aa in attempt_answers:
        chapter = aa.question.chapter
        if chapter.id not in chapter_stats:
            chapter_stats[chapter.id] = {
                "chapter_id": chapter.id,
                "chapter_name": chapter.name,
                "subject_name": chapter.subject.name,
                "total_answered": 0,
                "correct_count": 0,
            }
        chapter_stats[chapter.id]["total_answered"] += 1
        if aa.is_correct:
            chapter_stats[chapter.id]["correct_count"] += 1

    by_chapter = []
    for stat in chapter_stats.values():
        accuracy = round((stat["correct_count"] / stat["total_answered"]) * 100, 1) if stat["total_answered"] > 0 else 0
        by_chapter.append(schemas.ChapterStatResponse(**stat, accuracy=accuracy))

    return schemas.StudentStatsResponse(
        total_attempts=total_attempts,
        average_score=average_score,
        highest_score=highest_score,
        by_chapter=by_chapter,
    )


@router.get("/admin/stats/overview", response_model=schemas.AdminOverviewResponse)
def admin_overview(db: Session = Depends(get_db), admin: models.User = Depends(require_admin)):
    total_students = db.query(models.User).filter(models.User.role == "student").count()
    total_questions = db.query(models.Question).count()
    total_exams = db.query(models.Exam).count()
    attempts = db.query(models.Attempt).all()
    total_attempts = len(attempts)
    average_score = round(sum(a.score for a in attempts) / total_attempts, 2) if total_attempts > 0 else 0

    return schemas.AdminOverviewResponse(
        total_students=total_students,
        total_questions=total_questions,
        total_exams=total_exams,
        total_attempts=total_attempts,
        average_score=average_score,
    )


# ---------- TỔNG QUAN CHO ADMIN ----------

@router.get("/admin/subjects/overview", response_model=List[schemas.SubjectOverviewItem])
def admin_subjects_overview(db: Session = Depends(get_db), admin: models.User = Depends(require_admin)):
    subjects = db.query(models.Subject).all()
    result = []
    for s in subjects:
        chapter_ids = [c.id for c in db.query(models.Chapter).filter(models.Chapter.subject_id == s.id)]
        total_questions = (
            db.query(models.Question).filter(models.Question.chapter_id.in_(chapter_ids)).count()
            if chapter_ids else 0
        )
        result.append(schemas.SubjectOverviewItem(
            id=s.id, name=s.name, code=s.code,
            total_chapters=len(chapter_ids), total_questions=total_questions,
        ))
    return result


@router.get("/subjects/{subject_id}/overview", response_model=schemas.SubjectDetailOverview)
def subject_detail_overview(
    subject_id: int,
    db: Session = Depends(get_db),
    admin: models.User = Depends(require_admin),
):
    subject = db.query(models.Subject).filter(models.Subject.id == subject_id).first()
    if not subject:
        raise HTTPException(status_code=404, detail="Môn học không tồn tại")

    chapters = db.query(models.Chapter).filter(models.Chapter.subject_id == subject_id).all()
    chapter_data = []
    total_questions = 0
    for c in chapters:
        count = db.query(models.Question).filter(models.Question.chapter_id == c.id).count()
        total_questions += count
        chapter_data.append(schemas.ChapterWithCount(id=c.id, name=c.name, question_count=count))

    return schemas.SubjectDetailOverview(
        subject_id=subject.id, subject_name=subject.name,
        total_chapters=len(chapters), total_questions=total_questions,
        chapters=chapter_data,
    )